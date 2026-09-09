"""把 SPEC_AM_826X_83XX_uart_protocol_v1.8.xlsx 抽成可读的 markdown。

排版规律（列号固定）：
    [0]Index [2]名称 [38]Write/Read      <- 命令块起始行
    [2]TAG [6]Checksum [10]CmdID [14]DataLen [18]Data结构 [32]Comment
    [18]续行...                          <- Data 的详细说明

一个命令块 = 起始行 + TAG 行 + 若干续行，直到下一个起始行。
"""
import re
import sys

import openpyxl

XLSX = sys.argv[1] if len(sys.argv) > 1 else "SPEC_AM_826X_83XX_uart_protocol_v1.8.xlsx"
SHEETS = ["system", "Display", "Network", "audio", "misc", "customer",
          "keyCode", "Passthrough"]

C_IDX, C_NAME, C_TAG, C_CS, C_CMD, C_LEN, C_DATA, C_CMT, C_WR = \
    0, 2, 2, 6, 10, 14, 18, 32, 38


def cell(row, i):
    if i < len(row) and row[i] is not None:
        return str(row[i]).strip()
    return ""


def parse_sheet(ws):
    """返回 [{name, cmd, tag, dlen, wr, data:[...], comment:[...]}]"""
    rows = list(ws.iter_rows(values_only=True))
    out = []
    cur = None

    for r in rows:
        idx = cell(r, C_IDX)
        name = cell(r, C_NAME)
        cmd = cell(r, C_CMD)
        data = cell(r, C_DATA)
        cmt = cell(r, C_CMT)

        # 起始行：Index 是数字且第 2 列是名称（不是 TAG 值）
        if re.fullmatch(r"\d+", idx) and name and name not in (
                "GT", "PL", "LP", "LL", "PP", "TAG"):
            if cur:
                out.append(cur)
            cur = {"idx": idx, "name": name, "wr": cell(r, C_WR),
                   "cmd": "", "tag": "", "dlen": "",
                   "data": [], "comment": []}
            continue

        if cur is None:
            continue

        # TAG 行：带 Command ID
        if re.fullmatch(r"0x[0-9A-Fa-f]{4}", cmd):
            cur["cmd"] = cmd.lower().replace("0X", "0x")
            cur["tag"] = name if name in ("GT", "PL", "LP", "LL", "PP") else ""
            cur["dlen"] = cell(r, C_LEN)

        if data:
            cur["data"].append(data)
        if cmt:
            cur["comment"].append(cmt)

    if cur:
        out.append(cur)
    return [c for c in out if c["cmd"]]


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    all_cmds = {}

    print("# UART 协议 spec 提取（v1.8）\n")
    print("从 `SPEC_AM_826X_83XX_uart_protocol_v1.8.xlsx` 机械提取，"
          "便于和代码里的命令表对照。\n")

    for sh in SHEETS:
        if sh not in wb.sheetnames:
            continue
        cmds = parse_sheet(wb[sh])
        if not cmds:
            continue

        print(f"\n## {sh}（{len(cmds)} 条）\n")
        for c in cmds:
            wr = f" `{c['wr']}`" if c["wr"] else ""
            tag = f" TAG={c['tag']}" if c["tag"] else ""
            dl = f" len={c['dlen']}" if c["dlen"] else ""
            print(f"### {c['cmd']} {c['name']}{wr}{tag}{dl}\n")

            for d in c["data"]:
                for line in d.split("\n"):
                    line = line.strip()
                    if line:
                        print(f"    {line}")
            if c["comment"]:
                print()
                for m in c["comment"]:
                    for line in m.split("\n"):
                        line = line.strip()
                        if line:
                            print(f"> {line}")
            print()

            all_cmds.setdefault(c["cmd"], []).append((sh, c["name"]))

    print(f"\n## 命令 ID 汇总（{len(all_cmds)} 个）\n")
    for k in sorted(all_cmds):
        where = "; ".join(f"{s}/{n}" for s, n in all_cmds[k])
        print(f"- `{k}` {where}")


if __name__ == "__main__":
    main()
